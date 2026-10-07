#!/usr/bin/env python3
"""Genera CSVs de catálogo y red comercial reales para esekau.

schema_name = esekau
Precio = c/dto. + Impuestos. Una lista. is_mock=false.
Excluye POP precio < 1 y exhibidores Ferrero Pascua 11483-11486.
"""
from __future__ import annotations

import csv
import math
import re
from collections import defaultdict
from pathlib import Path

import openpyxl

SCHEMA = "esekau"
ROOT = Path(__file__).resolve().parents[1]
INPUTS = ROOT / "inputs"
OUT = ROOT / "outputs"

LISTA_XLSX = INPUTS / "Lista de Precios.xlsx"
ZONAS_XLSX = INPUTS / "Zona Agente (ex Aguilar).xlsx"

FERRERO_EXHIB = {11483, 11484, 11485, 11486}
AGUILAR_PHONE = "5493517000355"

FREQ_DIA = {"L": "lunes", "Ma": "martes", "Mi": "miercoles", "J": "jueves", "V": "viernes"}
RADIO_META = {
    "!_LICEO": {"nombre": "Liceo", "color": "#E74C3C", "codigo": "R355-LICEO"},
    "!_LOS PLATANOS": {"nombre": "Los Plátanos", "color": "#3498DB", "codigo": "R355-PLATANOS"},
    "!_ITUZAINGO": {"nombre": "Ituzaingó", "color": "#2ECC71", "codigo": "R355-ITUZAINGO"},
    "!_LA FRANCE/ALTO VERDE": {"nombre": "La France / Alto Verde", "color": "#F39C12", "codigo": "R355-LAFRANCE"},
    "!_TALLERES ESTE/PATRICIOS": {"nombre": "Talleres Este / Patricios", "color": "#9B59B6", "codigo": "R355-TALLERES"},
}

BRANDS = [
    "PANTENE", "HEAD & SHOULDERS", "HEAD AND SHOULDERS", "ORAL-B", "ORAL B",
    "PAMPERS", "ALWAYS", "DOWNY", "GILLETTE", "PRESTOBARBA", "VENUS",
    "OLD SPICE", "MAGISTRAL", "ZORRO", "PLUSBELLE", "ACE", "ARIEL",
    "POETT", "AYUDIN", "AYUDÍN", "TERMA", "DR LEMON", "FERNET",
    "BACARDI", "GANCIA", "TARAGUI", "TARAGÜI", "DURACELL", "PRIME",
    "BONAFIDE", "FELFORT", "FERRERO", "COCA-COLA", "COCA COLA",
]


def num(x) -> float:
    if x is None or x == "":
        return 0.0
    if isinstance(x, (int, float)):
        return float(x)
    try:
        return float(str(x).replace(",", ".").strip())
    except ValueError:
        return 0.0


def clean(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip()


def word_count(text: str) -> int:
    return len(re.findall(r"\S+", text))


def clip_words(text: str, lo: int = 10, hi: int = 25) -> str:
    words = re.findall(r"\S+", text)
    if len(words) > hi:
        text = " ".join(words[:hi]).rstrip(".,;") + "."
        words = re.findall(r"\S+", text)
    if len(words) < lo:
        text = text.rstrip(".") + " Lista mayorista ESEKA-U Córdoba."
    return text


def detect_marca(nombre: str) -> str:
    up = nombre.upper()
    for brand in BRANDS:
        if brand in up:
            return brand.title().replace("Dr Lemon", "Dr Lemon")
    token = re.split(r"[\s/]+", nombre)[0]
    return token.title() if token else ""


def descripcion(nombre: str, linea: str, rubro: str, caja: int, marca: str) -> str:
    noun = clean(rubro).lower() or "producto"
    noun = noun.split("(")[0].strip() or "producto"
    bits = [f"{noun.capitalize()} {nombre}"]
    if marca:
        bits.append(f"marca {marca}")
    if caja > 1:
        bits.append(f"caja x{caja}")
    if linea:
        bits.append(f"línea {linea.title()}")
    text = ", ".join(bits) + "."
    return clip_words(text)


def aliases(nombre: str, code: str, linea: str, rubro: str, marca: str, caja: int, barcode: str) -> str:
    parts = [nombre, code, linea, rubro, marca, "unidad", "caja", "bulto"]
    if caja > 1:
        parts += [f"caja x{caja}", f"bulto x{caja}", f"x{caja}"]
    if barcode:
        parts.append(barcode)
    seen: set[str] = set()
    out: list[str] = []
    for p in parts:
        p = clean(p)
        key = p.lower()
        if not p or key in seen:
            continue
        seen.add(key)
        out.append(p)
    return "|".join(out)


def rotacion(linea: str, idx: int, total_linea: int) -> float:
    base = {
        "PROCTER": 0.88,
        "DREAMCO": 0.72,
        "CLOROX": 0.68,
        "CEPAS": 0.62,
        "VARIOS": 0.42,
    }.get((linea or "").upper(), 0.4)
    spread = 0.12
    t = 0 if total_linea <= 1 else idx / (total_linea - 1)
    return round(max(0.1, min(0.95, base + spread * (0.5 - t))), 3)


def convex_hull(points: list[tuple[float, float]]) -> list[tuple[float, float]]:
    pts = sorted(set(points))
    if len(pts) <= 2:
        return pts

    def cross(o, a, b):
        return (a[0] - o[0]) * (b[1] - o[1]) - (a[1] - o[1]) * (b[0] - o[0])

    lower: list[tuple[float, float]] = []
    for p in pts:
        while len(lower) >= 2 and cross(lower[-2], lower[-1], p) <= 0:
            lower.pop()
        lower.append(p)
    upper: list[tuple[float, float]] = []
    for p in reversed(pts):
        while len(upper) >= 2 and cross(upper[-2], upper[-1], p) <= 0:
            upper.pop()
        upper.append(p)
    return lower[:-1] + upper[:-1]


def irregular_polygon(lats: list[float], lngs: list[float]) -> str:
    pts = list(zip(lngs, lats))  # lon, lat
    hull = convex_hull(pts)
    if len(hull) < 3:
        clat = sum(lats) / len(lats)
        clng = sum(lngs) / len(lngs)
        r = 0.004
        hull = [
            (clng + r * math.cos(math.radians(a)), clat + r * 0.85 * math.sin(math.radians(a)))
            for a in (0, 50, 110, 165, 220, 280, 330)
        ]
    clng = sum(p[0] for p in hull) / len(hull)
    clat = sum(p[1] for p in hull) / len(hull)
    scale = 1.12
    ring = [
        (clng + (lon - clng) * scale, clat + (lat - clat) * scale)
        for lon, lat in hull
    ]
    if ring[0] != ring[-1]:
        ring.append(ring[0])
    coords = ", ".join(f"{lon:.6f} {lat:.6f}" for lon, lat in ring)
    return f"SRID=4326;MULTIPOLYGON((({coords})))"


def generate_catalog() -> tuple[list[dict], list[dict]]:
    wb = openpyxl.load_workbook(LISTA_XLSX, read_only=True, data_only=True)
    ws = wb["Hoja1"]
    linea = rubro = None
    kept: list[dict] = []
    excluded: list[dict] = []
    for row in ws.iter_rows(min_row=6, values_only=True):
        desc = clean(row[0])
        art = row[1]
        if desc.upper().startswith("LINEA:"):
            linea = clean(row[1]) or desc
            continue
        if desc.upper().startswith("RUBRO:"):
            rubro = clean(row[1]) or desc
            continue
        if art is None or not desc:
            continue
        try:
            art_n = int(float(art))
        except (TypeError, ValueError):
            continue
        cdto = num(row[7])
        imp = num(row[8])
        precio = round(cdto + imp, 2)
        caja = int(num(row[3]) or 1)
        barcode = clean(row[2])
        motivo = ""
        if precio < 1:
            motivo = "precio_menor_1"
        elif art_n in FERRERO_EXHIB:
            motivo = "exhibidor_ferrero_pascua"
        rec = {
            "product_code": str(art_n),
            "nombre": desc,
            "precio": precio,
            "caja": caja,
            "linea": linea or "VARIOS",
            "rubro": rubro or "OTROS",
            "barcode": barcode,
        }
        if motivo:
            rec["motivo"] = motivo
            excluded.append(rec)
        else:
            kept.append(rec)
    wb.close()

    by_linea: dict[str, list[dict]] = defaultdict(list)
    for r in kept:
        by_linea[r["linea"]].append(r)
    for linea_name, group in by_linea.items():
        for i, r in enumerate(group):
            r["rotacion_index"] = rotacion(linea_name, i, len(group))
            r["mental_priority"] = round(r["rotacion_index"] * 0.9, 3)

    products = []
    for r in kept:
        marca = detect_marca(r["nombre"])
        products.append(
            {
                "product_code": r["product_code"],
                "nombre": r["nombre"],
                "precio_lista_1": f"{r['precio']:.2f}",
                "stock": str(int(10 + r["rotacion_index"] * 490)),
                "unidades_por_bulto": str(r["caja"]),
                "unidad_minima_de_venta": "unidad",
                "umv_tipo": "unidad",
                "categoria_1": r["linea"],
                "categoria_2": r["rubro"],
                "categoria_3": marca,
                "categoria_4": r["linea"],
                "aliases": aliases(
                    r["nombre"], r["product_code"], r["linea"], r["rubro"], marca, r["caja"], r["barcode"]
                ),
                "rotacion_index": str(r["rotacion_index"]),
                "mental_priority": str(r["mental_priority"]),
                "descripcion": descripcion(r["nombre"], r["linea"], r["rubro"], r["caja"], marca),
                "image_url": "",
                "en_catalogo": "true",
                "is_mock": "false",
                "fuente_hoja": "Lista de Precios.xlsx",
            }
        )
    return products, excluded


def normalize_phone(raw: str, codigo: int, seen: set[str]) -> tuple[str, str]:
    digits = re.sub(r"[^0-9]", "", raw or "")
    if digits.startswith("54") and len(digits) >= 11:
        phone = digits
    elif digits.startswith("9") and len(digits) >= 10:
        phone = "54" + digits
    elif len(digits) == 10 and digits.startswith("351"):
        phone = "549" + digits
    elif len(digits) >= 8:
        phone = "549" + digits[-10:] if len(digits) >= 10 else "549351" + digits[-8:]
    else:
        phone = ""
    if phone and len(phone) >= 11 and phone not in seen:
        seen.add(phone)
        return phone, "real"
    ph = f"5493518{codigo:07d}"
    while ph in seen:
        codigo += 100000
        ph = f"5493518{codigo:07d}"
    seen.add(ph)
    return ph, "placeholder"


def generate_red() -> tuple[list[dict], list[dict], list[dict]]:
    wb = openpyxl.load_workbook(ZONAS_XLSX, read_only=True, data_only=True)
    ws = wb["CLIENTES"]
    rows = list(ws.iter_rows(values_only=True))
    headers = [clean(h) for h in rows[0]]
    clients_raw = []
    for row in rows[1:]:
        d = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}
        clients_raw.append(d)
    wb.close()

    vendedores = [
        {
            "vendedor_codigo": "355",
            "nombre": "Joaquín Aguilar",
            "telefono": AGUILAR_PHONE,
            "email": "joaquin.aguilar@esekau.com",
            "activo": "true",
            "is_mock": "false",
            "zona": "Córdoba Capital — Agente",
            "codigo_ruta": "R355",
        }
    ]

    by_radio: dict[str, list[dict]] = defaultdict(list)
    for d in clients_raw:
        radio = clean(d.get("RADIO_LÍNEA_5"))
        by_radio[radio].append(d)

    zonas = []
    radio_order = []
    for radio, meta in RADIO_META.items():
        pts = []
        freq = None
        for d in by_radio.get(radio, []):
            try:
                pts.append((float(d["LATITUD"]), float(d["LONGITUD"])))
            except (TypeError, ValueError):
                continue
            freq = freq or clean(d.get("FRECUENCIA_LÍNEA_5"))
        if not pts:
            continue
        dia = FREQ_DIA.get(freq or "", "lunes")
        lats = [p[0] for p in pts]
        lngs = [p[1] for p in pts]
        zonas.append(
            {
                "zona_codigo": meta["codigo"],
                "nombre": meta["nombre"],
                "dia_visita": dia,
                "color_hex": meta["color"],
                "vendedor_codigo": "355",
                "geometry_wkt": irregular_polygon(lats, lngs),
                "is_mock": "false",
                "zone_type": "sales",
                "description": f"Radio {meta['nombre']} del agente Joaquín Aguilar. Córdoba Capital.",
                "n_clientes_origen": str(len(by_radio.get(radio, []))),
            }
        )
        radio_order.append(radio)

    seen_phones: set[str] = {AGUILAR_PHONE}
    clientes = []
    for d in clients_raw:
        radio = clean(d.get("RADIO_LÍNEA_5"))
        if radio not in RADIO_META:
            continue
        codigo = int(float(d["CÓD_CLIENTE"]))
        nombre = clean(d.get("NOMBRE"))
        direccion = clean(d.get("DIRECCIÓN"))
        freq = clean(d.get("FRECUENCIA_LÍNEA_5"))
        dia = FREQ_DIA.get(freq, "lunes")
        phone, phone_origen = normalize_phone(clean(d.get("CELULAR")), codigo, seen_phones)
        clientes.append(
            {
                "cliente_codigo": str(codigo),
                "razon_social": nombre,
                "nombre": nombre,
                "phone_number": phone,
                "phone_origen": phone_origen,
                "zona_codigo": RADIO_META[radio]["codigo"],
                "vendedor_codigo": "355",
                "lat": str(d["LATITUD"]),
                "lng": str(d["LONGITUD"]),
                "lista_precios_id": "1",
                "is_mock": "false",
                "dia_de_visita": dia,
                "direccion": direccion,
                "localidad": clean(d.get("LOCALIDAD")),
                "categoria_pg": clean(d.get("CATEGORÍA_PG")),
                "iva": clean(d.get("TIPO_RESPONSABLE_IVA")),
            }
        )
    return vendedores, zonas, clientes


def write_csv(path: Path, rows: list[dict], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def main() -> int:
    print(f"[*] schema_name={SCHEMA}")
    products, excluded = generate_catalog()
    vendedores, zonas, clientes = generate_red()

    prod_fields = [
        "product_code", "nombre", "precio_lista_1", "stock", "unidades_por_bulto",
        "unidad_minima_de_venta", "umv_tipo", "categoria_1", "categoria_2",
        "categoria_3", "categoria_4", "aliases", "rotacion_index", "mental_priority",
        "descripcion", "image_url", "en_catalogo", "is_mock", "fuente_hoja",
    ]
    write_csv(OUT / "phase-01-productos.csv", products, prod_fields)
    write_csv(
        OUT / "phase-01-lista-precios-1.csv",
        [{"product_code": p["product_code"], "precio_unidad": p["precio_lista_1"], "is_mock": "false"} for p in products],
        ["product_code", "precio_unidad", "is_mock"],
    )
    write_csv(
        OUT / "phase-01-listas-precios.csv",
        [{"lista_precios_id": "1", "nombre": "Lista Base Córdoba", "is_mock": "false", "es_publica": "true"}],
        ["lista_precios_id", "nombre", "is_mock", "es_publica"],
    )
    write_csv(
        OUT / "phase-01-excluidos-pop.csv",
        excluded,
        ["product_code", "nombre", "precio", "motivo", "linea", "rubro"],
    )
    write_csv(
        OUT / "phase-04-vendedores.csv",
        vendedores,
        ["vendedor_codigo", "nombre", "telefono", "email", "activo", "is_mock", "zona", "codigo_ruta"],
    )
    write_csv(
        OUT / "phase-04-zonas.csv",
        zonas,
        ["zona_codigo", "nombre", "dia_visita", "color_hex", "vendedor_codigo", "geometry_wkt", "is_mock", "zone_type", "description", "n_clientes_origen"],
    )
    write_csv(
        OUT / "phase-04-clientes.csv",
        clientes,
        [
            "cliente_codigo", "razon_social", "nombre", "phone_number", "phone_origen",
            "zona_codigo", "vendedor_codigo", "lat", "lng", "lista_precios_id",
            "is_mock", "dia_de_visita", "direccion", "localidad", "categoria_pg", "iva",
        ],
    )

    ph_real = sum(1 for c in clientes if c["phone_origen"] == "real")
    print(f"[OK] productos={len(products)} excluidos_pop={len(excluded)}")
    print(f"[OK] vendedores=1 zonas={len(zonas)} clientes={len(clientes)} tel_real={ph_real} tel_placeholder={len(clientes)-ph_real}")
    print(f"[OK] Aguilar telefono placeholder {AGUILAR_PHONE}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

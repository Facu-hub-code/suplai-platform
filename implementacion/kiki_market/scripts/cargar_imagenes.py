#!/usr/bin/env python3
"""Sube Foto 1 del Excel incrustado a Storage y actualiza kiki_market.productos.image_url.

schema_name = kiki_market. Solo escribe URL real (nunca placeholder).
"""
from __future__ import annotations

import asyncio
import csv
import os
import re
import sys
from pathlib import Path

import asyncpg
import openpyxl
import requests
from dotenv import load_dotenv

SCHEMA = "kiki_market"
HERE = Path(__file__).resolve()
ROOT = HERE.parents[3]
TENANT = HERE.parents[1]
EXCEL = TENANT / "inputs" / "KIKI_MARKET_catalogo_con_fotos_incrustadas.xlsx"
CSV_OUT = TENANT / "outputs" / "phase-01-imagenes.csv"
PRODUCTOS_CSV = TENANT / "outputs" / "phase-01-productos.csv"
BUCKET = f"products-{SCHEMA}"
COL_SKU = 1
COL_NOMBRE = 3
COL_FOTO1 = 5  # E
FILA_INICIO = 2
BATCH = 80


def force_pooler(url: str) -> str:
    return url.replace(":5432/", ":6543/")


def _load_envs() -> None:
    for path in (
        Path("/Users/facundolorenzo/Documents/SuplaiSales/source/backend-supabase/.env"),
        ROOT.parent / "backend-supabase" / ".env",
        ROOT / ".env",
    ):
        if path.exists():
            load_dotenv(path, override=False)
            print(f"[*] env: {path}")
    url = os.getenv("SUPABASE_DB_URL_POOLER") or os.getenv("SUPABASE_DB_URL") or ""
    url = force_pooler(url)
    if url:
        os.environ["SUPABASE_DB_URL"] = url
        os.environ["SUPABASE_DB_URL_POOLER"] = url


def norm_code(sku: str) -> str:
    s = str(sku or "").strip()
    if s.endswith(".0") and s[:-2].replace("-", "").isdigit():
        s = s[:-2]
    s = re.sub(r"\s+", "-", s)
    return s[:80]


def split_codes(raw: str) -> list[str]:
    parts = re.split(r"\s*,\s*", str(raw or "").strip())
    out = []
    for p in parts:
        p = p.lstrip("-").strip()
        c = norm_code(p)
        if c:
            out.append(c)
    return out


def safe_filename(code: str, ext: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9._-]+", "-", code).strip("-") or "sku"
    return f"{slug}.{ext}"


def foto1_by_row(ws) -> dict[int, object]:
    found: dict[int, object] = {}
    for img in ws._images:
        try:
            col = img.anchor._from.col + 1
            row = img.anchor._from.row + 1
        except Exception:
            continue
        if col == COL_FOTO1 and row not in found:
            found[row] = img
    return found


def ensure_bucket(supabase_url: str, key: str) -> None:
    headers = {
        "Authorization": f"Bearer {key}",
        "apikey": key,
        "Content-Type": "application/json",
    }
    resp = requests.post(
        f"{supabase_url.rstrip('/')}/storage/v1/bucket",
        headers=headers,
        json={"id": BUCKET, "name": BUCKET, "public": True},
        timeout=30,
    )
    if resp.status_code == 200:
        print(f"[*] bucket creado {BUCKET}")
    else:
        print(f"[*] bucket {BUCKET}: HTTP {resp.status_code}")


def upload_image(supabase_url: str, key: str, filename: str, image_bytes: bytes, ext: str) -> str | None:
    upload_url = f"{supabase_url.rstrip('/')}/storage/v1/object/{BUCKET}/{filename}"
    headers = {
        "Authorization": f"Bearer {key}",
        "apikey": key,
        "Content-Type": f"image/{ext}",
        "x-upsert": "true",
    }
    resp = requests.post(upload_url, headers=headers, data=image_bytes, timeout=60)
    if resp.status_code not in (200, 201):
        print(f"[WARN] upload {filename} HTTP {resp.status_code}: {resp.text[:180]}")
        return None
    return f"{supabase_url.rstrip('/')}/storage/v1/object/public/{BUCKET}/{filename}"


async def main() -> int:
    print(f"[*] schema_name confirmado: {SCHEMA}")
    _load_envs()
    db_url = force_pooler(os.getenv("SUPABASE_DB_URL_POOLER") or os.getenv("SUPABASE_DB_URL") or "")
    supabase_url = (os.getenv("SUPABASE_URL") or "").rstrip("/")
    key = os.getenv("SUPABASE_SERVICE_ROLE_KEY") or ""
    if not db_url or not supabase_url or not key:
        print("[FAIL] Falta SUPABASE_DB_URL / SUPABASE_URL / SUPABASE_SERVICE_ROLE_KEY", file=sys.stderr)
        return 1
    if not EXCEL.exists():
        print(f"[FAIL] No está el Excel: {EXCEL}", file=sys.stderr)
        return 1

    conn = await asyncpg.connect(db_url, statement_cache_size=0)
    try:
        await conn.execute("SET statement_timeout = 0")
        db_rows = await conn.fetch(f"SELECT product_code, nombre, image_url FROM {SCHEMA}.productos")
        db_codes = {r["product_code"] for r in db_rows}
        print(f"[*] productos en {SCHEMA}: {len(db_codes)}")

        print(f"[*] leyendo {EXCEL.name}...")
        wb = openpyxl.load_workbook(EXCEL, data_only=True)
        ws = wb.active
        fotos = foto1_by_row(ws)
        print(f"[*] hoja {ws.title} filas={ws.max_row} foto1={len(fotos)}")

        mappings: dict[str, str] = {}
        log_rows: list[dict[str, str]] = []
        uploaded = 0
        skipped_no_match = 0
        skipped_no_img = 0
        ensure_bucket(supabase_url, key)

        for r in range(FILA_INICIO, ws.max_row + 1):
            raw_sku = ws.cell(r, COL_SKU).value
            nombre = str(ws.cell(r, COL_NOMBRE).value or "").strip()
            codes = [c for c in split_codes(raw_sku) if c in db_codes]
            img = fotos.get(r)
            if not codes:
                skipped_no_match += 1
                continue
            if img is None:
                skipped_no_img += 1
                for code in codes:
                    log_rows.append(
                        {
                            "product_code": code,
                            "nombre": nombre,
                            "image_url": "",
                            "estado": "sin_foto1",
                        }
                    )
                continue

            fmt = (getattr(img, "format", None) or "png").lower()
            ext = "jpg" if fmt in {"jpeg", "jpg"} else ("png" if fmt == "png" else fmt)
            image_bytes = img.ref.getvalue()
            for code in codes:
                filename = safe_filename(code, ext)
                public_url = upload_image(supabase_url, key, filename, image_bytes, ext)
                if not public_url:
                    log_rows.append(
                        {
                            "product_code": code,
                            "nombre": nombre,
                            "image_url": "",
                            "estado": "error_upload",
                        }
                    )
                    continue
                mappings[code] = public_url
                uploaded += 1
                log_rows.append(
                    {
                        "product_code": code,
                        "nombre": nombre,
                        "image_url": public_url,
                        "estado": "ok",
                    }
                )
            if uploaded and uploaded % 50 == 0:
                print(f"    subidas {uploaded}...")

        print(f"[*] upload ok={len(mappings)} skip_sku={skipped_no_match} skip_img={skipped_no_img}")
        if not mappings:
            print("[FAIL] Ninguna imagen asociada.", file=sys.stderr)
            return 1

        updates = [(url, code) for code, url in mappings.items()]
        print(f"[*] UPDATE {SCHEMA}.productos.image_url ({len(updates)} filas)...")
        for i in range(0, len(updates), BATCH):
            chunk = updates[i : i + BATCH]
            await conn.executemany(
                f"""
                UPDATE {SCHEMA}.productos
                SET image_url = $1, updated_at = now()
                WHERE product_code = $2
                """,
                chunk,
            )
            print(f"    db {i + 1}–{i + len(chunk)}")

        verify = await conn.fetchrow(
            f"""
            SELECT
              COUNT(*) FILTER (WHERE image_url IS NOT NULL AND btrim(image_url) <> '') AS con_foto,
              COUNT(*) AS total
            FROM {SCHEMA}.productos
            """
        )
        print(f"[VERIFY] con_foto={verify['con_foto']} total={verify['total']}")
    finally:
        await conn.close()

    CSV_OUT.parent.mkdir(exist_ok=True)
    with CSV_OUT.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["product_code", "nombre", "image_url", "estado"])
        w.writeheader()
        w.writerows(log_rows)

    if PRODUCTOS_CSV.exists():
        rows = list(csv.DictReader(PRODUCTOS_CSV.open(encoding="utf-8")))
        fields = list(rows[0].keys()) if rows else []
        for row in rows:
            url = mappings.get(row.get("product_code", ""))
            if url:
                row["image_url"] = url
        with PRODUCTOS_CSV.open("w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, fieldnames=fields)
            w.writeheader()
            w.writerows(rows)
        print(f"[*] sync {PRODUCTOS_CSV.name}")

    backend_url = os.getenv("BACKEND_URL", "https://web-production-f544f.up.railway.app").rstrip("/")
    codes = list(mappings)
    try:
        resp = requests.post(f"{backend_url}/{SCHEMA}/productos/vectorize", json=codes, timeout=120)
        print(f"[*] vectorize HTTP {resp.status_code}: {resp.text[:200]}")
    except Exception as e:
        print(f"[WARN] vectorize: {e}")

    print(f"[SUCCESS] {len(mappings)} fotos en {SCHEMA}")
    return 0


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))

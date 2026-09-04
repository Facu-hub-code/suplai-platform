#!/usr/bin/env python3
"""Sube fotos reales de PRODUCTOS-GRAL.xlsx + PNGs de combos helados.

No escribe placeholders. No pisa image_url si la subida falla.
schema_name = cordoba_frost
"""
from __future__ import annotations

import asyncio
import csv
import os
import sys
from collections import defaultdict
from pathlib import Path

import asyncpg
import requests
from dotenv import load_dotenv
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
PLATFORM = ROOT.parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from generar_fase01 import (  # noqa: E402
    EXISTING_SKIP,
    SKIP_ZERO_PRICE,
    load_rows,
    next_generated_code,
    prefix_for,
    sanitize_code,
)

EXCEL = ROOT / "inputs" / "PRODUCTOS-GRAL.xlsx"
CSV_PATH = ROOT / "outputs" / "phase-01-productos-gral.csv"
COMBO_PNGS = {
    "COM-HEL-INICIAL": ROOT / "inputs/combos-helados-agosto-2026/combo-inicial.png",
    "COM-HEL-MEDIO": ROOT / "inputs/combos-helados-agosto-2026/combo-medio.png",
    "COM-HEL-PREMIUM": ROOT / "inputs/combos-helados-agosto-2026/combo-premium.png",
}
SCHEMA = "cordoba_frost"
BUCKET = "products-cordoba_frost"
REPORT = ROOT / "outputs" / "phase-01-fotos-reporte.csv"


def excel_row_to_code() -> dict[int, str]:
    raw_rows = load_rows()
    used = {"COM-HEL-INICIAL", "COM-HEL-MEDIO", "COM-HEL-PREMIUM", "ENVIO-DOM"}
    used.update(EXISTING_SKIP)
    seq = [20000]
    sanitized = []
    counts: dict[str, int] = defaultdict(int)
    for r in raw_rows:
        code, nota = sanitize_code(r["codigo_raw"], r["marca"])
        sanitized.append({**r, "code": code, "nota": nota})
        if code:
            counts[code] += 1
    seen_once: set[str] = set()
    mapping: dict[int, str] = {}
    for r in sanitized:
        precio = r["precio"]
        if precio is None or precio <= 0:
            continue
        code = r["code"]
        nota = r["nota"]
        if not code:
            code = next_generated_code(prefix_for(r["marca"], r["cat"]), used, seq)
            nota = "generado_sin_codigo_excel"
        elif code in SKIP_ZERO_PRICE:
            continue
        elif counts[code] > 1:
            if code in seen_once:
                code = next_generated_code(prefix_for(r["marca"], r["cat"]), used, seq)
                nota = "duplicado"
            else:
                seen_once.add(code)
        if code in used and nota.startswith("sanitizado"):
            code = next_generated_code(prefix_for(r["marca"], r["cat"]), used, seq)
        used.add(code)
        mapping[r["excel_row"]] = code
    return mapping


def first_image_per_row(ws) -> dict[int, object]:
    by_row: dict[int, object] = {}
    for img in ws._images:
        r = img.anchor._from.row + 1
        if r not in by_row:
            by_row[r] = img
    return by_row


def image_bytes(img) -> tuple[bytes, str]:
    fmt = (getattr(img, "format", None) or "png").lower()
    ext = "jpg" if fmt in ("jpeg", "jpg") else ("png" if fmt == "png" else fmt)
    if ext not in ("png", "jpg", "webp", "gif"):
        ext = "png"
    if hasattr(img, "_data") and callable(img._data):
        data = img._data()
    elif hasattr(img, "ref") and hasattr(img.ref, "getvalue"):
        data = img.ref.getvalue()
    else:
        raise RuntimeError("no pude extraer bytes de la imagen embebida")
    return data, ext


def ensure_bucket(supabase_url: str, key: str) -> None:
    url = f"{supabase_url.rstrip('/')}/storage/v1/bucket"
    headers = {"Authorization": f"Bearer {key}", "apikey": key, "Content-Type": "application/json"}
    r = requests.post(url, headers=headers, json={"id": BUCKET, "name": BUCKET, "public": True}, timeout=30)
    if r.status_code in (200, 201):
        print(f"✅ Bucket {BUCKET} creado")
    elif r.status_code in (400, 409) or "already exists" in (r.text or "").lower():
        print(f"[*] Bucket {BUCKET} ya existe")
    else:
        print(f"[WARN] bucket HTTP {r.status_code}: {r.text[:200]}")


def upload_bytes(supabase_url: str, key: str, filename: str, data: bytes, ext: str) -> str | None:
    upload_url = f"{supabase_url.rstrip('/')}/storage/v1/object/{BUCKET}/{filename}"
    mime = "image/jpeg" if ext == "jpg" else f"image/{ext}"
    headers = {
        "Authorization": f"Bearer {key}",
        "apikey": key,
        "Content-Type": mime,
        "x-upsert": "true",
    }
    resp = requests.post(upload_url, headers=headers, data=data, timeout=60)
    if resp.status_code not in (200, 201):
        print(f"[WARN] upload {filename} HTTP {resp.status_code}: {resp.text[:180]}")
        return None
    return f"{supabase_url.rstrip('/')}/storage/v1/object/public/{BUCKET}/{filename}"


async def main() -> int:
    load_dotenv(PLATFORM.parent / "backend-supabase/.env")
    load_dotenv(PLATFORM / ".env", override=False)
    db_url = os.getenv("SUPABASE_DB_URL")
    supabase_url = os.getenv("SUPABASE_URL")
    supabase_key = os.getenv("SUPABASE_SERVICE_ROLE_KEY")
    if not db_url or not supabase_url or not supabase_key:
        print("[FAIL] Faltan SUPABASE_DB_URL / SUPABASE_URL / SUPABASE_SERVICE_ROLE_KEY")
        return 1

    row_codes = excel_row_to_code()
    print(f"[*] Filas Excel mapeadas a SKU: {len(row_codes)}")

    wb = load_workbook(EXCEL, data_only=True)
    ws = wb.active
    imgs = first_image_per_row(ws)
    print(f"[*] Filas con foto: {len(imgs)}")

    ensure_bucket(supabase_url, supabase_key)

    mappings: dict[str, str] = {}
    report = []
    skipped_no_sku = 0
    skipped_no_img = 0

    for excel_row, code in sorted(row_codes.items()):
        img = imgs.get(excel_row)
        if not img:
            skipped_no_img += 1
            report.append({"product_code": code, "excel_row": excel_row, "fuente": "excel", "ok": "sin_foto"})
            continue
        try:
            data, ext = image_bytes(img)
        except Exception as e:
            print(f"[WARN] no pude leer imagen fila {excel_row} {code}: {e}")
            report.append({"product_code": code, "excel_row": excel_row, "fuente": "excel", "ok": f"error_bytes:{e}"})
            continue
        filename = f"{code}.{ext}"
        public_url = upload_bytes(supabase_url, supabase_key, filename, data, ext)
        if public_url:
            mappings[code] = public_url
            report.append({"product_code": code, "excel_row": excel_row, "fuente": "excel", "ok": "ok", "url": public_url})
        else:
            report.append({"product_code": code, "excel_row": excel_row, "fuente": "excel", "ok": "upload_fail"})

    for extra_row in set(imgs) - set(row_codes):
        skipped_no_sku += 1
        report.append({"product_code": "", "excel_row": extra_row, "fuente": "excel", "ok": "fila_sin_sku"})

    for code, png in COMBO_PNGS.items():
        if code in mappings:
            continue
        if not png.exists():
            print(f"[WARN] falta PNG {png}")
            continue
        public_url = upload_bytes(supabase_url, supabase_key, f"{code}.png", png.read_bytes(), "png")
        if public_url:
            mappings[code] = public_url
            report.append({"product_code": code, "excel_row": "", "fuente": "png_combo", "ok": "ok", "url": public_url})

    wb.close()
    print(f"[*] Subidas OK: {len(mappings)} | filas sin foto: {skipped_no_img} | fotos sin SKU: {skipped_no_sku}")

    exist_set: set[str] = set()
    conn = await asyncpg.connect(db_url, statement_cache_size=0)
    try:
        existing = await conn.fetch(
            f"SELECT product_code FROM {SCHEMA}.productos WHERE product_code = ANY($1::text[])",
            list(mappings),
        )
        exist_set = {r["product_code"] for r in existing}
        missing = [c for c in mappings if c not in exist_set]
        if missing:
            print(f"[WARN] SKUs con foto que no están en BD: {missing}")
        updates = [(url, code) for code, url in mappings.items() if code in exist_set]
        if updates:
            await conn.executemany(
                f"UPDATE {SCHEMA}.productos SET image_url = $1, updated_at = NOW() WHERE product_code = $2 AND (image_url IS NULL OR btrim(image_url) = '')",
                updates,
            )
        n_now = await conn.fetchval(
            f"SELECT COUNT(*) FROM {SCHEMA}.productos WHERE image_url IS NOT NULL AND btrim(image_url) <> ''"
        )
        print(f"✅ BD actualizada. Productos con foto ahora: {n_now}")
    finally:
        await conn.close()

    with REPORT.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["product_code", "excel_row", "fuente", "ok", "url"])
        w.writeheader()
        for row in report:
            w.writerow({k: row.get(k, "") for k in w.fieldnames})
    print(f"[*] Reporte: {REPORT}")

    if CSV_PATH.exists() and mappings:
        rows = []
        with CSV_PATH.open(encoding="utf-8") as f:
            reader = csv.DictReader(f)
            fields = reader.fieldnames
            for row in reader:
                code = row.get("product_code")
                if code in mappings:
                    row["image_url"] = mappings[code]
                rows.append(row)
        with CSV_PATH.open("w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, fieldnames=fields)
            w.writeheader()
            w.writerows(rows)
        print("✅ CSV phase-01-productos-gral.csv sincronizado")

    codes = [c for c in mappings if c in exist_set]
    if codes:
        backend = os.getenv("BACKEND_URL", "https://web-production-f544f.up.railway.app").rstrip("/")
        for i in range(0, len(codes), 100):
            chunk = codes[i : i + 100]
            try:
                resp = requests.post(f"{backend}/{SCHEMA}/productos/vectorize", json=chunk, timeout=60)
                print(f"[*] vectorize {len(chunk)} HTTP {resp.status_code}")
            except Exception as e:
                print(f"[WARN] vectorize: {e}")
    return 0


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))

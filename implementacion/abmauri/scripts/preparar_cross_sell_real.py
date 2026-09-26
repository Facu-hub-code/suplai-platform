#!/usr/bin/env python3
"""Deriva cross-sell real desde comprobantes AB Mauri 2024–2026.

No escribe en Supabase. Deduplica SKUs dentro de cada comprobante, calcula
soporte/confianza/lift y selecciona 15 reglas direccionales diversificadas.
"""
from __future__ import annotations

import csv
import json
import math
from collections import Counter, defaultdict
from datetime import datetime
from itertools import combinations
from pathlib import Path

from openpyxl import load_workbook

EXPECTED = 15
MIN_SUPPORT = 25
ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "inputs" / "2026 - ABMauri - Resumen 23_09.xlsx"
PRODUCTS = ROOT / "outputs" / "phase-01-productos.csv"
TAXONOMY = ROOT / "outputs" / "phase-01-1-propuesta-categorias.json"
CROSS_OUT = ROOT / "outputs" / "phase-03-cross-sell.csv"
UP_OUT = ROOT / "outputs" / "phase-03-up-sell.csv"
AUDIT_OUT = ROOT / "outputs" / "phase-03-cross-sell-auditoria.csv"

CROSS_FIELDS = ["base_product_code", "related_product_code", "reason", "is_mock"]
AUDIT_FIELDS = [
    "priority",
    "base_product_code",
    "base_nombre",
    "base_familia",
    "related_product_code",
    "related_nombre",
    "related_familia",
    "support_comprobantes",
    "confidence",
    "lift",
    "score",
]


def code(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value or "").strip()


def main() -> None:
    with PRODUCTS.open(encoding="utf-8-sig", newline="") as handle:
        products = {row["product_code"]: row for row in csv.DictReader(handle)}
    taxonomy_data = json.loads(TAXONOMY.read_text(encoding="utf-8"))
    taxonomy = {
        item["product_code"]: item["tags"] for item in taxonomy_data["products"]
    }

    workbook = load_workbook(INPUT, read_only=True, data_only=True)
    baskets: dict[tuple[str, str, str, str], set[str]] = defaultdict(set)
    for index, row in enumerate(
        workbook["PEDIDOS HISTORICOS AL 23092026"].iter_rows(values_only=True)
    ):
        if (
            index == 0
            or not isinstance(row[2], datetime)
            or row[2].year < 2024
            or not row[6]
        ):
            continue
        sku = code(row[6])
        if sku not in products:
            continue
        key = (
            str(row[3]),
            str(row[4]),
            str(row[0]),
            row[2].date().isoformat(),
        )
        baskets[key].add(sku)
    workbook.close()

    item_support: Counter[str] = Counter()
    pair_support: Counter[tuple[str, str]] = Counter()
    for basket in baskets.values():
        for sku in basket:
            item_support[sku] += 1
        for left, right in combinations(sorted(basket), 2):
            pair_support[(left, right)] += 1

    basket_count = len(baskets)
    rules: list[dict[str, object]] = []
    for (left, right), support in pair_support.items():
        if support < MIN_SUPPORT:
            continue
        if taxonomy[left]["2"] == taxonomy[right]["2"]:
            continue
        lift = support * basket_count / (item_support[left] * item_support[right])
        for base, related in ((left, right), (right, left)):
            confidence = support / item_support[base]
            score = confidence * math.log1p(support) * min(lift, 5)
            rules.append(
                {
                    "base": base,
                    "related": related,
                    "support": support,
                    "confidence": confidence,
                    "lift": lift,
                    "score": score,
                    "transition": (taxonomy[base]["2"], taxonomy[related]["2"]),
                }
            )
    rules.sort(
        key=lambda rule: (
            -float(rule["score"]),
            -int(rule["support"]),
            str(rule["base"]),
            str(rule["related"]),
        )
    )

    selected: list[dict[str, object]] = []
    base_count: Counter[str] = Counter()
    related_count: Counter[str] = Counter()
    transition_count: Counter[tuple[str, str]] = Counter()
    unordered_pairs: set[tuple[str, str]] = set()
    for rule in rules:
        base = str(rule["base"])
        related = str(rule["related"])
        pair_key = tuple(sorted((base, related)))
        transition = rule["transition"]
        if (
            pair_key in unordered_pairs
            or base_count[base] >= 1
            or related_count[related] >= 3
            or transition_count[transition] >= 3
        ):
            continue
        selected.append(rule)
        unordered_pairs.add(pair_key)
        base_count[base] += 1
        related_count[related] += 1
        transition_count[transition] += 1
        if len(selected) == EXPECTED:
            break

    if len(selected) != EXPECTED:
        raise RuntimeError(f"Solo se pudieron seleccionar {len(selected)}/{EXPECTED} reglas")

    cross_rows: list[dict[str, object]] = []
    audit_rows: list[dict[str, object]] = []
    for priority, rule in enumerate(selected, 1):
        base = str(rule["base"])
        related = str(rule["related"])
        support = int(rule["support"])
        confidence = float(rule["confidence"])
        lift = float(rule["lift"])
        reason = (
            f"Comprados juntos en {support} comprobantes entre 2024 y 2026; "
            f"{confidence:.1%} de los comprobantes del producto base incluyeron "
            f"el sugerido; lift {lift:.2f}."
        )
        cross_rows.append(
            {
                "base_product_code": base,
                "related_product_code": related,
                "reason": reason,
                "is_mock": False,
            }
        )
        audit_rows.append(
            {
                "priority": priority,
                "base_product_code": base,
                "base_nombre": products[base]["nombre"],
                "base_familia": taxonomy[base]["2"],
                "related_product_code": related,
                "related_nombre": products[related]["nombre"],
                "related_familia": taxonomy[related]["2"],
                "support_comprobantes": support,
                "confidence": round(confidence, 6),
                "lift": round(lift, 6),
                "score": round(float(rule["score"]), 6),
            }
        )

    with CROSS_OUT.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=CROSS_FIELDS)
        writer.writeheader()
        writer.writerows(cross_rows)
    with UP_OUT.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=CROSS_FIELDS)
        writer.writeheader()
    with AUDIT_OUT.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=AUDIT_FIELDS)
        writer.writeheader()
        writer.writerows(audit_rows)

    multi = [basket for basket in baskets.values() if len(basket) >= 2]
    print(f"comprobantes={basket_count} multisku={len(multi)}")
    print(f"cross_sell={len(cross_rows)} up_sell=0")
    print(
        f"support_min={min(int(rule['support']) for rule in selected)} "
        f"support_max={max(int(rule['support']) for rule in selected)}"
    )


if __name__ == "__main__":
    main()

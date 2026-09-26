#!/usr/bin/env python3
"""Prepara pedidos piloto desde facturas reales AB Mauri de 2026.

Conserva cliente, fecha y composición real del comprobante. Como el origen no
incluye cantidades ni importes por línea, deduplica cada SKU y asigna cantidad
1, valuando con la lista piloto del cliente. Todo queda marcado is_mock=true.
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "inputs" / "2026 - ABMauri - Resumen 23_09.xlsx"
PRODUCTS = ROOT / "outputs" / "phase-01-productos.csv"
CLIENTS = ROOT / "outputs" / "phase-04-clientes.csv"
PRICES = ROOT / "outputs" / "phase-04-precios-piloto.csv"
ORDERS_OUT = ROOT / "outputs" / "phase-06-pedidos.csv"
ITEMS_OUT = ROOT / "outputs" / "phase-06-items-pedido.csv"
REPORT_OUT = ROOT / "outputs" / "phase-06-pilot-report.json"

ORDER_FIELDS = [
    "pedido_ref",
    "source_comprobante",
    "source_factura",
    "cliente_codigo",
    "source_cliente_codigo",
    "cliente_phone",
    "cliente_razon_social",
    "lista_precios_id",
    "fecha",
    "estado",
    "total",
    "notas",
    "is_mock",
    "es_pedido_abierto",
]
ITEM_FIELDS = [
    "pedido_ref",
    "cliente_phone",
    "product_code",
    "nombre",
    "cantidad_solicitada",
    "precio_unitario",
    "lista_precios_id",
    "fecha_pedido",
    "notas",
    "promo_aplicada",
    "is_mock",
    "source_row_count",
]


def clean_code(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value or "").strip()


def clean_ref(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value or "").strip()


def reference(prefix: str, key: tuple[str, str, str, str]) -> str:
    readable = re.sub(r"[^A-Za-z0-9]+", "-", key[1] or key[2]).strip("-")[:35]
    digest = hashlib.sha1("|".join(key).encode("utf-8")).hexdigest()[:10]
    return f"{prefix}-{readable or 'SINREF'}-{digest}"


def write_csv(path: Path, fields: list[str], rows: list[dict[str, object]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    with PRODUCTS.open(encoding="utf-8-sig", newline="") as handle:
        products = {
            row["product_code"]: row for row in csv.DictReader(handle)
        }
    with CLIENTS.open(encoding="utf-8", newline="") as handle:
        clients = {
            row["source_client_code"]: row for row in csv.DictReader(handle)
        }
    with PRICES.open(encoding="utf-8", newline="") as handle:
        prices = {
            (row["product_code"], int(row["lista_precios_id"])): Decimal(
                row["precio_unidad"]
            )
            for row in csv.DictReader(handle)
        }

    workbook = load_workbook(INPUT, read_only=True, data_only=True)
    invoices: dict[tuple[str, str, str, str], Counter[str]] = defaultdict(Counter)
    for index, row in enumerate(
        workbook["PEDIDOS HISTORICOS AL 23092026"].iter_rows(values_only=True)
    ):
        if (
            index == 0
            or not isinstance(row[2], datetime)
            or row[2].year != 2026
            or not row[6]
        ):
            continue
        client_code = clean_code(row[0])
        key = (
            client_code,
            clean_ref(row[3]),
            clean_ref(row[4]),
            row[2].date().isoformat(),
        )
        invoices[key][clean_code(row[6])] += 1
    workbook.close()

    usable = [
        (key, sku_counts)
        for key, sku_counts in invoices.items()
        if key[0] in clients
        and sku_counts
        and all(sku in products for sku in sku_counts)
    ]
    usable.sort(key=lambda pair: (pair[0][3], pair[0][0], pair[0][1], pair[0][2]))

    order_rows: list[dict[str, object]] = []
    item_rows: list[dict[str, object]] = []
    for key, sku_counts in usable:
        source_client_code, comprobante, factura, date = key
        client = clients[source_client_code]
        list_id = int(client["lista_precios_id"])
        pedido_ref = reference("ERP2026", key)
        total = Decimal("0")
        for sku in sorted(sku_counts):
            price = prices[(sku, list_id)]
            total += price
            item_rows.append(
                {
                    "pedido_ref": pedido_ref,
                    "cliente_phone": client["phone_number"],
                    "product_code": sku,
                    "nombre": products[sku]["nombre"],
                    "cantidad_solicitada": 1,
                    "precio_unitario": str(price),
                    "lista_precios_id": list_id,
                    "fecha_pedido": date,
                    "notas": (
                        "Cantidad piloto=1; el origen contiene "
                        f"{sku_counts[sku]} fila(s) para este SKU, sin cantidad explícita."
                    ),
                    "promo_aplicada": False,
                    "is_mock": True,
                    "source_row_count": sku_counts[sku],
                }
            )
        order_rows.append(
            {
                "pedido_ref": pedido_ref,
                "source_comprobante": comprobante,
                "source_factura": factura,
                "cliente_codigo": client["client_code"],
                "source_cliente_codigo": source_client_code,
                "cliente_phone": client["phone_number"],
                "cliente_razon_social": client["razon_social"],
                "lista_precios_id": list_id,
                "fecha": date,
                "estado": "facturado",
                "total": str(total),
                "notas": (
                    "Factura real 2026; composición real por SKU. "
                    "Cantidades normalizadas a 1 y precios estimados del piloto."
                ),
                "is_mock": True,
                "es_pedido_abierto": False,
            }
        )

    open_candidates = [
        (key, sku_counts)
        for key, sku_counts in reversed(usable)
        if len(sku_counts) >= 2
    ]
    selected_open: list[tuple[tuple[str, str, str, str], Counter[str]]] = []
    seen_clients: set[str] = set()
    for key, sku_counts in open_candidates:
        if key[0] in seen_clients:
            continue
        selected_open.append((key, sku_counts))
        seen_clients.add(key[0])
        if len(selected_open) == 7:
            break
    if len(selected_open) != 7:
        raise RuntimeError(f"Solo se encontraron {len(selected_open)}/7 pedidos abiertos")

    open_date = "2026-09-26"
    for number, (key, sku_counts) in enumerate(selected_open, 1):
        source_client_code = key[0]
        client = clients[source_client_code]
        list_id = int(client["lista_precios_id"])
        pedido_ref = f"PILOTO-OPEN-{number:03d}"
        total = Decimal("0")
        selected_skus = sorted(sku_counts)[:4]
        for sku in selected_skus:
            price = prices[(sku, list_id)]
            total += price
            item_rows.append(
                {
                    "pedido_ref": pedido_ref,
                    "cliente_phone": client["phone_number"],
                    "product_code": sku,
                    "nombre": products[sku]["nombre"],
                    "cantidad_solicitada": 1,
                    "precio_unitario": str(price),
                    "lista_precios_id": list_id,
                    "fecha_pedido": open_date,
                    "notas": "Pedido abierto del piloto; cantidad sintética=1.",
                    "promo_aplicada": False,
                    "is_mock": True,
                    "source_row_count": sku_counts[sku],
                }
            )
        order_rows.append(
            {
                "pedido_ref": pedido_ref,
                "source_comprobante": key[1],
                "source_factura": key[2],
                "cliente_codigo": client["client_code"],
                "source_cliente_codigo": source_client_code,
                "cliente_phone": client["phone_number"],
                "cliente_razon_social": client["razon_social"],
                "lista_precios_id": list_id,
                "fecha": open_date,
                "estado": "pendiente",
                "total": str(total),
                "notas": (
                    "Pedido abierto sintético para probar la bandeja; "
                    "composición basada en una factura real."
                ),
                "is_mock": True,
                "es_pedido_abierto": True,
            }
        )

    write_csv(ORDERS_OUT, ORDER_FIELDS, order_rows)
    write_csv(ITEMS_OUT, ITEM_FIELDS, item_rows)
    report = {
        "schema": "abmauri",
        "source_year": 2026,
        "source_invoices": len(invoices),
        "closed_orders": len(usable),
        "open_orders": len(selected_open),
        "orders_total": len(order_rows),
        "items_total": len(item_rows),
        "quantity_rule": "1 por SKU único en comprobante",
        "pricing_rule": "lista piloto asignada al cliente",
        "is_mock": True,
        "excluded_invoices": len(invoices) - len(usable),
    }
    REPORT_OUT.write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

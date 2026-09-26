#!/usr/bin/env python3
"""Prepara red comercial completa para el piloto interno AB Mauri.

Usa los 16.554 clientes reales del Excel. Solo inventa teléfonos, vendedores y
asignaciones de listas, todos marcados como datos de piloto.
"""
from __future__ import annotations

import csv
import hashlib
import json
from collections import Counter
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "inputs" / "2026 - ABMauri - Resumen 23_09.xlsx"
PRICE_INPUT = ROOT / "outputs" / "phase-01-lista-precios-1.csv"
OUT = ROOT / "outputs"

SELLER_FIELDS = ["seller_ref", "nombre", "telefono", "email", "activo", "is_mock"]
CLIENT_FIELDS = [
    "client_code",
    "source_client_code",
    "razon_social",
    "nombre_fantasia",
    "cuit",
    "email",
    "direccion",
    "calle",
    "localidad",
    "provincia",
    "phone_number",
    "seller_ref",
    "seller_nombre",
    "lista_precios_id",
    "lista_precios_nombre",
    "whatsapp_estado",
    "is_mock",
    "metadata_json",
]
LIST_FIELDS = ["id", "nombre", "descripcion", "multiplicador", "is_mock"]
PRICE_FIELDS = ["product_code", "lista_precios_id", "precio_unidad", "is_mock"]

LISTS = [
    {
        "id": 1,
        "nombre": "Piloto interno - Default",
        "descripcion": "Precio estimado base del piloto interno",
        "multiplicador": Decimal("1.00"),
        "is_mock": True,
    },
    {
        "id": 2,
        "nombre": "Piloto interno - Minorista",
        "descripcion": "Precio estimado minorista: 15% sobre lista default",
        "multiplicador": Decimal("1.15"),
        "is_mock": True,
    },
    {
        "id": 3,
        "nombre": "Piloto interno - Mayorista",
        "descripcion": "Precio estimado mayorista: 10% bajo lista default",
        "multiplicador": Decimal("0.90"),
        "is_mock": True,
    },
]


def clean_code(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value or "").strip()


def clean_optional(value: object) -> str:
    text = str(value or "").strip()
    return "" if text.lower() in {"", "none", "null", "nan"} else text


def clean_cuit(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean_optional(value).replace("-", "").replace(".", "")


def deterministic_order(code: str) -> str:
    return hashlib.sha256(f"abmauri:{code}".encode("utf-8")).hexdigest()


def write_csv(path: Path, fields: list[str], rows: list[dict[str, object]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    sellers = [
        {
            "seller_ref": f"VP-{number:02d}",
            "nombre": f"Vendedor Piloto {number:02d}",
            "telefono": str(5491109000000 + number),
            "email": f"vendedor{number:02d}@suplaisales.mock",
            "activo": True,
            "is_mock": True,
        }
        for number in range(1, 11)
    ]

    workbook = load_workbook(INPUT, read_only=True, data_only=True)
    source_clients: list[dict[str, str]] = []
    for index, row in enumerate(
        workbook["CLIENTES AL 23092026"].iter_rows(values_only=True)
    ):
        if index == 0 or not any(row):
            continue
        source_code = clean_code(row[0])
        source_clients.append(
            {
                "client_code": source_code,
                "source_client_code": source_code,
                "razon_social": clean_optional(row[1]) or clean_optional(row[2]),
                "nombre_fantasia": clean_optional(row[2]) or clean_optional(row[1]),
                "cuit": clean_cuit(row[3]),
                "calle": clean_optional(row[4]),
                "localidad": clean_optional(row[5]),
                "provincia": clean_optional(row[6]),
                "direccion": clean_optional(row[7]),
                "email": clean_optional(row[8]).lower(),
            }
        )
    workbook.close()

    used_numeric_codes = {
        int(row["client_code"])
        for row in source_clients
        if row["client_code"].isdigit()
    }
    surrogate = 99_000_000
    for row in source_clients:
        if row["client_code"].isdigit():
            continue
        while surrogate in used_numeric_codes:
            surrogate += 1
        row["client_code"] = str(surrogate)
        used_numeric_codes.add(surrogate)
        surrogate += 1

    source_clients.sort(
        key=lambda row: deterministic_order(row["source_client_code"])
    )
    total = len(source_clients)
    if total != 16554:
        raise RuntimeError(f"Se esperaban 16554 clientes y se encontraron {total}")
    default_count = 9934
    minorista_count = 3310
    mayorista_count = 3310
    if default_count + minorista_count + mayorista_count != total:
        raise RuntimeError("La distribución de listas no suma el total de clientes")

    client_rows: list[dict[str, object]] = []
    for index, source in enumerate(source_clients):
        seller = sellers[index % len(sellers)]
        if index < default_count:
            price_list = LISTS[0]
        elif index < default_count + minorista_count:
            price_list = LISTS[1]
        else:
            price_list = LISTS[2]
        metadata = {
            "source": "ABMauri Excel CLIENTES AL 23092026",
            "source_client_code": source["source_client_code"],
            "pilot_internal": True,
            "real_fields": [
                "source_client_code",
                "razon_social",
                "nombre_fantasia",
                "cuit",
                "email",
                "direccion",
                "localidad",
                "provincia",
            ],
            "synthetic_fields": [
                "client_code cuando el código ERP original no es numérico",
                "phone_number",
                "seller_ref",
                "lista_precios_id",
            ],
        }
        client_rows.append(
            {
                **source,
                "phone_number": str(5491100000001 + index),
                "seller_ref": seller["seller_ref"],
                "seller_nombre": seller["nombre"],
                "lista_precios_id": price_list["id"],
                "lista_precios_nombre": price_list["nombre"],
                "whatsapp_estado": "no_validado",
                "is_mock": True,
                "metadata_json": json.dumps(metadata, ensure_ascii=False),
            }
        )

    with PRICE_INPUT.open(encoding="utf-8-sig", newline="") as handle:
        base_prices = list(csv.DictReader(handle))
    if len(base_prices) != 182:
        raise RuntimeError(f"Se esperaban 182 precios base y hay {len(base_prices)}")

    price_rows: list[dict[str, object]] = []
    for row in base_prices:
        base_price = Decimal(row["precio_unidad"])
        for price_list in LISTS:
            price_rows.append(
                {
                    "product_code": row["product_code"],
                    "lista_precios_id": price_list["id"],
                    "precio_unidad": str(
                        (base_price * price_list["multiplicador"]).quantize(
                            Decimal("0.01")
                        )
                    ),
                    "is_mock": True,
                }
            )

    write_csv(OUT / "phase-04-vendedores.csv", SELLER_FIELDS, sellers)
    write_csv(OUT / "phase-04-clientes.csv", CLIENT_FIELDS, client_rows)
    write_csv(OUT / "phase-04-listas-precios-piloto.csv", LIST_FIELDS, LISTS)
    write_csv(OUT / "phase-04-precios-piloto.csv", PRICE_FIELDS, price_rows)
    write_csv(
        OUT / "phase-04-zonas.csv",
        ["zone_ref", "nombre", "notas"],
        [],
    )

    report = {
        "schema": "abmauri",
        "clientes": len(client_rows),
        "vendedores": len(sellers),
        "zonas": 0,
        "telefonos_sinteticos": len({row["phone_number"] for row in client_rows}),
        "codigos_erp_no_numericos_remapeados": sum(
            row["client_code"] != row["source_client_code"] for row in client_rows
        ),
        "clientes_por_vendedor": dict(
            Counter(str(row["seller_ref"]) for row in client_rows)
        ),
        "clientes_por_lista": dict(
            Counter(str(row["lista_precios_id"]) for row in client_rows)
        ),
        "listas": len(LISTS),
        "precios": len(price_rows),
        "prices_per_list": {
            str(list_id): count
            for list_id, count in Counter(
                str(row["lista_precios_id"]) for row in price_rows
            ).items()
        },
        "synthetic_fields": [
            "clients.phone_number",
            "clients.seller_ref",
            "clients.lista_precios_id",
            "vendedores.*",
            "listas_precios 1-3",
            "precios_productos listas 1-3",
        ],
    }
    (OUT / "phase-04-pilot-report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

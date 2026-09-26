#!/usr/bin/env python3
"""Genera la Fase 1 del piloto interno AB Mauri.

Los productos y su actividad provienen del Excel real. Los precios son
referencias web extrapoladas por familia; el fallback determinístico queda
identificado como inventado. No escribe en Supabase.
"""
from __future__ import annotations

import csv
import hashlib
import math
import re
import unicodedata
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "inputs" / "2026 - ABMauri - Resumen 23_09.xlsx"
OUT = ROOT / "outputs"
FULL = ROOT / "inputs" / "catalogo-completo.csv"

PRODUCT_FIELDS = [
    "product_code",
    "nombre",
    "precio_lista_1",
    "stock",
    "unidades_por_bulto",
    "unidad_minima_de_venta",
    "umv_tipo",
    "categoria_1",
    "categoria_2",
    "categoria_3",
    "categoria_4",
    "aliases",
    "rotacion_index",
    "mental_priority",
    "descripcion",
    "image_url",
    "en_catalogo",
    "is_mock",
    "fuente_hoja",
]

PRICE_FIELDS = [
    "lista_precios_id",
    "nombre",
    "multiplicador_sobre_lista_1",
    "product_code",
    "precio_unidad",
    "is_mock",
]

AUDIT_FIELDS = [
    "product_code",
    "nombre",
    "familia_precio",
    "presentacion_total",
    "precio_unidad",
    "metodo_precio",
    "fuente_url",
]

WEB_SOURCES = {
    "levadura_fresca": "https://distribuidoraalpes.com.ar/producto/levadura-prensada-virgen-x-500-grs-calsa/",
    "levadura_seca": "https://tienda.elnuevoemporio.com.ar/tienda/otros-alimentos/levadura-seca-instantanea-mauripan-x-500gr",
    "margarina": "https://distribuidoraalpes.com.ar/categoria-producto/grasas-margarinas/",
    "grasa": "https://distribuidorachichi.com.ar/productos/77-grasa-pastelgras-x-20-kg.html",
    "premezcla": "https://www.distribuidorajumbalay.com.ar/levadura-premezcla-pan-de-queso-x-2-kg-calsa--det--PRM000",
    "mejorador": "https://sapore.com.ar/s-0101610-mejorador-formula-1-calsa-x-10-kg/",
    "crema": "https://marinosrl.com.ar/producto/400_crema-pastelera-calsa-(en-frio)-x-3-kg",
    "esencia": "https://distribuidorachichi.com.ar/productos/73-0100100_ESENCIA-DE-VAINILLA-CALSA-X-5-L-7790813000016.html",
    "kolaroma": "https://distribuidorachichi.com.ar/productos/73-0100100_ESENCIA-DE-VAINILLA-CALSA-X-5-L-7790813000016.html",
    "extracto_malta": "https://distribuidorachichi.com.ar/productos/74-0101710_EXTRACTO-DE-MALTA-CALSA-X-28-K.html",
}

# Precio por kg/litro derivado de las referencias web del 25/09/2026.
WEB_RATE = {
    "levadura_fresca": 11_400,
    "levadura_seca": 24_640,
    "margarina": 7_240,
    "grasa": 6_450,
    "premezcla": 10_536,
    "mejorador": 8_200,
    "crema": 11_059,
    "esencia": 3_000,
    "kolaroma": 26_600,
    "extracto_malta": 6_800,
}

DIRECT_WEB_PRICE = {
    "11016001": 5_700,
    "11016082": 108_300,
    "15017001": 12_320,
    "15017082": 234_080,
    "73166030": 21_072.61,
    "73166143": 120_113.88,
    "63250042": 82_000,
    "63250144": 155_800,
    "40154114": 144_780,
    "32025093": 128_999.92,
    "72810145": 189_114.60,
    "72401129": 15_000.01,
    "72401155": 57_000.04,
}

UMV = {
    "CS": "caja",
    "BAG": "bolsa",
    "BAL": "balde",
    "PIL": "pilon",
    "PA1": "paquete",
    "BI1": "bidon",
    "POT": "pote",
    "EA": "unidad",
    "DSP": "display",
    "DR": "tambor",
    "ES1": "estuche",
    "BLS": "bolsa",
    "PAN": "pan",
    "SOB": "sobre",
    "BOT": "botella",
}


def clean_code(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value or "").strip()


def clean_text(value: object) -> str:
    return " ".join(str(value or "").split())


def fold(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value or "")
    return "".join(c for c in normalized if not unicodedata.combining(c)).upper()


def family(product: dict[str, object]) -> str:
    text = fold(
        " ".join(
            [
                str(product["nombre"]),
                str(product["division"]),
                str(product["rubro"]),
                str(product["segmento"]),
            ]
        )
    )
    if "KOLAROMA" in text:
        return "kolaroma"
    if "EXTRACTO" in text and "MALTA" in text:
        return "extracto_malta"
    if str(product["division"]).lower() == "levadura":
        return "levadura_seca" if "SECA" in text or "INSTANT" in text else "levadura_fresca"
    if "MARGAR" in text or "MELANGE" in text:
        return "margarina"
    if (
        str(product["division"]).lower() in {"grasos", "productos grasos"}
        or "GRASA" in text
        or "SHORTENING" in text
        or "OLEOMARGARINA" in text
    ):
        return "grasa"
    if "PREMEZCLA" in text or "PREM." in text:
        return "premezcla"
    if "MEJORADOR" in text or re.search(r"\bMEJ\b", text):
        return "mejorador"
    if "CREMA" in text:
        return "crema"
    if "ESENCIA" in text or "COLORANTE" in text:
        return "esencia"
    if "CHOCOLATE" in text or "BANO" in text:
        return "chocolate"
    if "POLVO" in text or "ADITIVO" in text:
        return "aditivo"
    return "otros"


def quantity_total(nombre: str, um_code: str) -> tuple[float, str, int]:
    text = fold(nombre).replace(",", ".")

    display = re.search(r"(\d+)D/(\d+)(?:FP|S|P)/(\d+(?:\.\d+)?)(?:G)?\b", text)
    if display:
        units = int(display.group(1)) * int(display.group(2))
        return units * float(display.group(3)) / 1000, "kg", units

    envelopes = re.search(r"CJ\s*(\d+)E/BOL\s*(\d+(?:\.\d+)?)(?:G)?\b", text)
    if envelopes:
        units = int(envelopes.group(1))
        return units * float(envelopes.group(2)) / 1000, "kg", units

    box = re.search(
        r"CJ\s*(\d+)\s*(?:PIL|PAQ|BOL|D)?\s*X\s*(\d+(?:\.\d+)?)\s*(KG|G|LT|L)\b",
        text,
    )
    if box:
        units = int(box.group(1))
        amount = float(box.group(2))
        unit = box.group(3)
        total = units * amount / 1000 if unit == "G" else units * amount
        return total, "l" if unit in {"L", "LT"} else "kg", units

    box_packages = re.search(
        r"CJ\s*(\d+)\s*(?:PAQ|PIL|BOL|SOB)\s*(\d+(?:\.\d+)?)\s*(KG|G|LT|L)\b",
        text,
    )
    if box_packages:
        units = int(box_packages.group(1))
        amount = float(box_packages.group(2))
        unit = box_packages.group(3)
        total = units * amount / 1000 if unit == "G" else units * amount
        return total, "l" if unit in {"L", "LT"} else "kg", units

    box_total = re.search(r"CJ\s*X?\s*(\d+(?:\.\d+)?)\s*(KG|G|LT|L)\b", text)
    if box_total:
        amount = float(box_total.group(1))
        unit = box_total.group(2)
        total = amount / 1000 if unit == "G" else amount
        return total, "l" if unit in {"L", "LT"} else "kg", 1

    single = re.search(
        r"(?:BOL|BOLSA|BIDON|PAQ|POTE|BALDE|TAMBOR)?\s*(\d+(?:\.\d+)?)\s*(KG|G|LT|L)\b",
        text,
    )
    if single:
        amount = float(single.group(1))
        unit = single.group(2)
        total = amount / 1000 if unit == "G" else amount
        return total, "l" if unit in {"L", "LT"} else "kg", 1

    defaults = {
        "CS": (20.0, "kg"),
        "BAG": (10.0, "kg"),
        "BAL": (10.0, "kg"),
        "PIL": (5.0, "kg"),
        "PA1": (0.5, "kg"),
        "BI1": (5.0, "l"),
        "POT": (1.0, "kg"),
    }
    amount, unit = defaults.get(um_code, (1.0, "kg"))
    return amount, unit, 1


def deterministic_factor(code: str) -> float:
    digest = hashlib.sha256(code.encode("utf-8")).digest()
    return 0.94 + (int.from_bytes(digest[:2], "big") / 65535) * 0.12


def round_commercial(value: float) -> int:
    return max(1_000, int(round(value / 100.0) * 100))


def estimate_price(product: dict[str, object]) -> tuple[int, str, str, str]:
    code = str(product["sku"])
    fam = family(product)
    amount, unit, units = quantity_total(str(product["nombre"]), str(product["um"]))
    presentation = f"{amount:g} {unit}"

    if code in DIRECT_WEB_PRICE:
        return (
            round_commercial(DIRECT_WEB_PRICE[code]),
            fam,
            presentation,
            "web_directo",
        )

    rate = WEB_RATE.get(fam)
    if rate:
        bulk_discount = 0.95 if units > 1 else 1.0
        value = rate * amount * bulk_discount * deterministic_factor(code)
        return round_commercial(value), fam, presentation, "web_extrapolado"

    fallback_rate = {
        "chocolate": 12_000,
        "aditivo": 9_000,
        "otros": 10_000,
    }.get(fam, 10_000)
    value = fallback_rate * amount * deterministic_factor(code)
    return round_commercial(value), fam, presentation, "formula_inventada"


def aliases(nombre: str, rubro: str, segmento: str, umv: str) -> str:
    cleaned = re.sub(r"\b(?:LP|I360|100A)\b", " ", fold(nombre))
    cleaned = " ".join(cleaned.split()).lower()
    values = [cleaned, clean_text(rubro).lower(), clean_text(segmento).lower(), f"{cleaned} por {umv}"]
    return "|".join(dict.fromkeys(v for v in values if v))


def description(product: dict[str, object], presentation: str) -> str:
    rubro = clean_text(product["rubro"]) or "Ingrediente"
    segmento = clean_text(product["segmento"])
    detail = segmento if segmento and segmento.lower() != rubro.lower() else clean_text(product["nombre"])
    return f"{rubro} Calsa para panificación, línea {detail}, presentación comercial de {presentation}."


def write_csv(path: Path, fields: list[str], rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    workbook = load_workbook(INPUT, read_only=True, data_only=True)
    products: dict[str, dict[str, object]] = {}
    product_sheet = workbook["LISTA PRODUCTOS AL 23092026"]
    for index, row in enumerate(product_sheet.iter_rows(values_only=True)):
        if index == 0 or not row[0]:
            continue
        sku = clean_code(row[1] or row[0])
        products[sku] = {
            "sku": sku,
            "nombre": clean_text(row[2]),
            "um": clean_text(row[3]),
            "um_desc": clean_text(row[4]),
            "division": clean_text(row[5]),
            "rubro": clean_text(row[6]),
            "segmento": clean_text(row[7]),
        }

    sold: Counter[str] = Counter()
    order_sheet = workbook["PEDIDOS HISTORICOS AL 23092026"]
    for index, row in enumerate(order_sheet.iter_rows(values_only=True)):
        if index == 0 or not row[6] or not isinstance(row[2], datetime) or row[2].year < 2024:
            continue
        sold[clean_code(row[6])] += 1
    workbook.close()

    active_codes = [
        code
        for code, product in products.items()
        if code in sold
        and str(product["division"]).lower() not in {"participaciones publicitarias", "servicios"}
        and str(product["rubro"]).lower() != "discontinuados"
    ]
    active_codes.sort(key=lambda code: (-sold[code], code))
    max_sales = sold[active_codes[0]]

    full_rows = [
        {
            "product_code": code,
            "nombre": product["nombre"],
            "unidad_medida": product["um"],
            "unidad_medida_desc": product["um_desc"],
            "division": product["division"],
            "rubro": product["rubro"],
            "segmento": product["segmento"],
            "lineas_venta_2024_2026": sold.get(code, 0),
        }
        for code, product in products.items()
    ]
    write_csv(
        FULL,
        [
            "product_code",
            "nombre",
            "unidad_medida",
            "unidad_medida_desc",
            "division",
            "rubro",
            "segmento",
            "lineas_venta_2024_2026",
        ],
        full_rows,
    )

    product_rows: list[dict[str, object]] = []
    price_rows: list[dict[str, object]] = []
    audit_rows: list[dict[str, object]] = []
    for position, code in enumerate(active_codes):
        product = products[code]
        price, fam, presentation, method = estimate_price(product)
        _, _, units = quantity_total(str(product["nombre"]), str(product["um"]))
        umv = UMV.get(str(product["um"]), "unidad")
        rotation = round(0.15 + 0.80 * math.log1p(sold[code]) / math.log1p(max_sales), 3)
        percentile = position / len(active_codes)
        priority = 0.9 if percentile < 0.2 else 0.6 if percentile < 0.5 else 0.3
        stock = max(10, int(round((20 + 480 * rotation) / 10) * 10))
        product_rows.append(
            {
                "product_code": code,
                "nombre": product["nombre"],
                "precio_lista_1": price,
                "stock": stock,
                "unidades_por_bulto": units,
                "unidad_minima_de_venta": umv,
                "umv_tipo": "display" if umv == "display" else "unidad",
                "categoria_1": product["division"],
                "categoria_2": product["rubro"],
                "categoria_3": product["segmento"],
                "categoria_4": umv.title(),
                "aliases": aliases(
                    str(product["nombre"]),
                    str(product["rubro"]),
                    str(product["segmento"]),
                    umv,
                ),
                "rotacion_index": rotation,
                "mental_priority": priority,
                "descripcion": description(product, presentation),
                "image_url": "",
                "en_catalogo": True,
                "is_mock": False,
                "fuente_hoja": "LISTA PRODUCTOS AL 23092026 + ventas 2024-2026",
            }
        )
        price_rows.append(
            {
                "lista_precios_id": 1,
                "nombre": "Piloto interno - precios estimados",
                "multiplicador_sobre_lista_1": 1.0,
                "product_code": code,
                "precio_unidad": price,
                "is_mock": True,
            }
        )
        audit_rows.append(
            {
                "product_code": code,
                "nombre": product["nombre"],
                "familia_precio": fam,
                "presentacion_total": presentation,
                "precio_unidad": price,
                "metodo_precio": method,
                "fuente_url": WEB_SOURCES.get(fam, ""),
            }
        )

    assert len(product_rows) == 182
    assert len({row["product_code"] for row in product_rows}) == 182
    assert all(row["is_mock"] is False for row in product_rows)
    assert all(row["precio_unidad"] > 0 and row["is_mock"] is True for row in price_rows)
    assert all(
        10 <= len(str(row["descripcion"]).rstrip(".").split()) <= 25 for row in product_rows
    )

    write_csv(OUT / "phase-01-productos.csv", PRODUCT_FIELDS, product_rows)
    write_csv(OUT / "phase-01-lista-precios-1.csv", PRICE_FIELDS, price_rows)
    write_csv(OUT / "phase-01-precios-auditoria.csv", AUDIT_FIELDS, audit_rows)

    methods = Counter(str(row["metodo_precio"]) for row in audit_rows)
    print(f"productos={len(product_rows)} universo={len(full_rows)}")
    print("metodos_precio=" + ",".join(f"{key}:{value}" for key, value in sorted(methods.items())))
    print(f"precio_min={min(row['precio_unidad'] for row in price_rows)}")
    print(f"precio_max={max(row['precio_unidad'] for row in price_rows)}")


if __name__ == "__main__":
    main()
